from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from PIL import Image

from app.services.ocr_service import (
    OCRProcessingError,
    extract_text_from_pil_image,
)


SUPPORTED_DOCUMENT_EXTENSIONS = {".txt", ".pdf", ".docx"}


class DocumentProcessingError(Exception):
    pass


def extract_txt(file_bytes: bytes) -> tuple[str, dict]:
    encodings = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]

    for encoding in encodings:
        try:
            text = file_bytes.decode(encoding)

            return text.strip(), {
                "encoding": encoding,
            }

        except UnicodeDecodeError:
            continue

    raise DocumentProcessingError(
        "The TXT file encoding could not be detected."
    )


def extract_pdf(file_bytes: bytes) -> tuple[str, dict]:
    page_results: list[str] = []
    ocr_page_count = 0
    maximum_ocr_pages = 25

    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
            page_count = pdf.page_count

            for page_index, page in enumerate(pdf):
                page_number = page_index + 1
                page_text = page.get_text("text").strip()

                if page_text:
                    page_results.append(
                        f"--- Page {page_number} ---\n{page_text}"
                    )
                    continue

                if page_number > maximum_ocr_pages:
                    continue

                try:
                    matrix = fitz.Matrix(2.0, 2.0)

                    pixmap = page.get_pixmap(
                        matrix=matrix,
                        alpha=False,
                    )

                    image = Image.frombytes(
                        "RGB",
                        (pixmap.width, pixmap.height),
                        pixmap.samples,
                    )

                    ocr_text = extract_text_from_pil_image(image)

                    if ocr_text:
                        page_results.append(
                            f"--- Page {page_number} (OCR) ---\n"
                            f"{ocr_text}"
                        )
                        ocr_page_count += 1

                except OCRProcessingError:
                    continue

            metadata = {
                "page_count": page_count,
                "pages_with_text": len(page_results),
                "ocr_page_count": ocr_page_count,
                "maximum_ocr_pages": maximum_ocr_pages,
                "pdf_metadata": {
                    key: value
                    for key, value in (pdf.metadata or {}).items()
                    if value
                },
            }

    except Exception as exc:
        raise DocumentProcessingError(
            "The PDF could not be opened or is damaged."
        ) from exc

    return "\n\n".join(page_results).strip(), metadata


def extract_docx(file_bytes: bytes) -> tuple[str, dict]:
    try:
        document = Document(BytesIO(file_bytes))

    except Exception as exc:
        raise DocumentProcessingError(
            "The DOCX file could not be opened or is damaged."
        ) from exc

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    table_text: list[str] = []

    for table_number, table in enumerate(
        document.tables,
        start=1,
    ):
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if values:
                table_text.append(
                    f"Table {table_number}: "
                    + " | ".join(values)
                )

    text_parts = paragraphs + table_text

    metadata = {
        "paragraph_count": len(paragraphs),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
    }

    return "\n".join(text_parts).strip(), metadata


def process_document(
    file_name: str,
    file_bytes: bytes,
) -> dict:
    extension = Path(file_name).suffix.lower()

    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise DocumentProcessingError(
            f"Unsupported document type: {extension}"
        )

    if extension == ".txt":
        extracted_text, metadata = extract_txt(file_bytes)

    elif extension == ".pdf":
        extracted_text, metadata = extract_pdf(file_bytes)

    else:
        extracted_text, metadata = extract_docx(file_bytes)

    warnings: list[str] = []

    if not extracted_text:
        warnings.append(
            "No readable text was found. OCR may not have detected "
            "text, or the document may contain unsupported content."
        )

    return {
        "extracted_text": extracted_text,
        "metadata": metadata,
        "warnings": warnings,
    }